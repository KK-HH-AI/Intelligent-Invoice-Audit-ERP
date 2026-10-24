"""随单附件文本提取工具：PDF / Word / Excel → 文本，图片 → OCR。

由旧 api/doc_analysis.py 拆出"提取"半边：
- 只负责把文件变成文本（供 T6 audit_service 拼 DeepSeek 提示词）；
- 临时文件落盘、DeepSeek 调用等编排不在本文件：新流程中附件由 api 层
  存入 uploads/，审查时按路径直接提取；
- 提取失败沿用旧行为返回 "[...失败原因]" 字符串，由上层决定如何处理。
"""
import logging
import os

import cv2
import fitz  # PyMuPDF
import numpy as np
from PyPDF2 import PdfReader

from services.ocr_engine import get_ocr

logger = logging.getLogger(__name__)

# docx / pandas 为可选依赖：函数内懒加载，缺库时走降级提示而不是拖垮整个模块


def extract_file_content(file_path: str, filename: str) -> str:
    """按扩展名分派的提取入口。

    Args:
        file_path: 本地文件路径
        filename: 原始文件名（用于判断文件类型）

    Returns:
        str: 提取的文本内容；不支持的类型抛 ValueError
    """
    file_ext = os.path.splitext(filename)[1].lower()
    if file_ext == ".pdf":
        return extract_pdf_content(file_path)
    if file_ext == ".docx":
        return extract_docx_content(file_path)
    if file_ext in (".xls", ".xlsx"):
        return extract_excel_content(file_path)
    raise ValueError(f"不支持的文件类型: {file_ext}")


def extract_pdf_content(file_path: str) -> str:
    """PDF 内容提取器（支持文本型 PDF 和扫描版 PDF）。

    处理流程：
        1. 优先使用 PyPDF2 提取文本内容
        2. 如果文本为空，按页渲染成图像并做增强后走 OCR（扫描版）
    """
    try:
        with open(file_path, "rb") as f:
            reader = PdfReader(f)
        text = "\n".join(page.extract_text() or "" for page in reader.pages)

        # 检测到扫描版 PDF（文本为空）时启用 OCR 流程
        if not text.strip():
            doc = fitz.Document(file_path)
            full_text = []

            for page_num in range(len(doc)):
                try:
                    page = doc[page_num]
                    # 根据页面大小动态调整缩放比例，确保图像清晰度
                    zoom = 2.0
                    if max(page.rect.width, page.rect.height) > 2000:
                        zoom = 1.5
                    elif max(page.rect.width, page.rect.height) < 1000:
                        zoom = 2.5

                    pix = page.get_pixmap(
                        matrix=fitz.Matrix(zoom, zoom),
                        alpha=False,
                        colorspace=fitz.csRGB,
                    )
                    img_array = np.frombuffer(pix.samples, dtype=np.uint8)
                    img_array = img_array.reshape(pix.h, pix.w, pix.n)

                    # 统一图像格式
                    if pix.n == 4:
                        img_array = cv2.cvtColor(img_array, cv2.COLOR_RGBA2RGB)
                    elif pix.n == 1:
                        img_array = cv2.cvtColor(img_array, cv2.COLOR_GRAY2RGB)

                    # 1. 灰度化 2. 自适应阈值增强对比度 3. 降噪 4. 锐化
                    gray = cv2.cvtColor(img_array, cv2.COLOR_RGB2GRAY)
                    binary = cv2.adaptiveThreshold(
                        gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
                        cv2.THRESH_BINARY, 11, 2,
                    )
                    denoised = cv2.fastNlMeansDenoising(binary, None, 10, 7, 21)
                    kernel = np.array([[-1, -1, -1], [-1, 9, -1], [-1, -1, -1]])
                    sharpened = cv2.filter2D(denoised, -1, kernel)

                    # 按对比度选择最优结果：处理显著提升则用处理图，否则轻度增强原图
                    orig_contrast = cv2.Laplacian(gray, cv2.CV_64F).var()
                    proc_contrast = cv2.Laplacian(sharpened, cv2.CV_64F).var()
                    if proc_contrast > orig_contrast * 1.2:
                        img_array = cv2.cvtColor(sharpened, cv2.COLOR_GRAY2RGB)
                    else:
                        img_array = cv2.convertScaleAbs(img_array, alpha=1.3, beta=10)
                        img_array = cv2.filter2D(img_array, -1, kernel)

                    ocr_result = process_image_ocr(img_array)
                    if ocr_result:
                        full_text.append(str(ocr_result))
                    else:
                        full_text.append(f"[第{page_num + 1}页OCR结果为空]")
                except Exception as e:
                    logger.error("PDF页面%d处理失败: %s", page_num + 1, e)
                    full_text.append(f"[第{page_num + 1}页处理失败: {e}]")
                finally:
                    # 释放内存
                    if "img_array" in locals():
                        del img_array
                    if "pix" in locals():
                        del pix

            if not full_text:
                text = "[未能提取到有效文本]"
            else:
                text = "\n".join(full_text)
            doc.close()
    except Exception as e:
        logger.error("PDF处理失败: %s, 错误: %s", file_path, e)
        return f"[PDF处理失败: {e}]"

    return text


def extract_docx_content(file_path: str) -> str:
    """使用 python-docx 提取 .docx：段落 + 表格（表格转 Markdown）。"""
    try:
        from docx import Document

        doc = Document(file_path)
        text = []

        for para in doc.paragraphs:
            if not para.text.strip():
                continue
            if para.style.name.startswith("Heading"):
                text.append(f"\n{para.text}\n")
            else:
                text.append(para.text)

        for table in doc.tables:
            text.append("\n")
            if table.rows:
                header = [cell.text.strip() for cell in table.rows[0].cells]
                text.append("| " + " | ".join(header) + " |")
                text.append("| " + " | ".join(["---" for _ in header]) + " |")
                for row in table.rows[1:]:
                    cells = [cell.text.strip() for cell in row.cells]
                    text.append("| " + " | ".join(cells) + " |")
            text.append("\n")

        return "\n".join(text).strip()
    except ImportError:
        logger.error("python-docx库未安装")
        return "[请安装python-docx库]"
    except Exception as e:
        logger.exception("DOCX解析失败: %s", e)
        return f"[DOCX提取失败: {e}]"


def extract_excel_content(file_path: str) -> str:
    """提取 Excel（.xls/.xlsx）为 Markdown 表格，保留工作表结构。"""
    try:
        import pandas as pd

        excel_data = pd.ExcelFile(file_path)
        sheet_contents = []

        for sheet_name in excel_data.sheet_names:
            df = pd.read_excel(excel_data, sheet_name=sheet_name)
            if df.empty:
                continue
            sheet_contents.append(f"\n## 工作表: {sheet_name}\n")
            md_table = df.to_markdown(index=False)
            if md_table:
                sheet_contents.append(md_table)
            sheet_contents.append("\n" + "-" * 80 + "\n")

        excel_data.close()

        if not sheet_contents:
            return "[Excel文件未包含有效数据]"
        return "\n".join(sheet_contents)
    except ImportError as e:
        logger.error("Excel处理需要pandas库: %s", e)
        return f"[Excel处理需要安装pandas库: {e}]"
    except Exception as e:
        logger.exception("Excel解析失败: %s", e)
        return f"[Excel提取失败: {e}]"


def process_image_ocr(image_data):
    """单张图片 OCR：返回 [{text, x, y}, ...]（文本框中心点坐标）。

    Args:
        image_data: numpy.ndarray，OpenCV 格式图像

    失败时返回错误信息字符串（旧行为，保持兼容）。
    """
    try:
        result = get_ocr().ocr(image_data)

        ocr_results = []
        for line in result[0]:
            text = line[1][0]
            position = line[0]

            x_coords = [p[0] for p in position]
            y_coords = [p[1] for p in position]
            x_mid = (min(x_coords) + max(x_coords)) / 2
            y_mid = (min(y_coords) + max(y_coords)) / 2

            ocr_results.append({"text": text, "x": int(x_mid), "y": int(y_mid)})
        return ocr_results
    except Exception as e:
        return f"[图片OCR处理失败: {e}]"
